"""
generate_test_cases.py
Generates AspireAI_Test_Cases.docx containing all test cases across 12 modules.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# ---------------------------------------------------------------------------
# Test-case data
# ---------------------------------------------------------------------------

MODULES = [
    {
        "id": "MODULE 1",
        "name": "User Authentication",
        "cases": [
            {
                "id": "TC-AUTH-001",
                "title": "Successful User Registration",
                "module": "User Authentication",
                "preconditions": "The application server is running. No user with the test email exists in the database.",
                "steps": (
                    "1. Send POST /api/auth/register with valid payload.\n"
                    "2. Verify the HTTP response code.\n"
                    "3. Verify the response body structure.\n"
                    "4. Confirm password is not returned in the response."
                ),
                "data": (
                    'POST /api/auth/register\n'
                    '{\n'
                    '  "firstName": "John",\n'
                    '  "lastName": "Doe",\n'
                    '  "email": "john.doe@example.com",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 201 Created.\n"
                    "Response body: { success: true, message: 'User registered successfully', data: { user: {...}, token: '...' } }.\n"
                    "data.user contains firstName, lastName, email.\n"
                    "data.user does NOT contain password.\n"
                    "data.token is a non-empty string."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-AUTH-002",
                "title": "Registration with Duplicate Email",
                "module": "User Authentication",
                "preconditions": "A user with email 'john.doe@example.com' already exists in the database.",
                "steps": (
                    "1. Send POST /api/auth/register with an email that already exists.\n"
                    "2. Verify the HTTP response code.\n"
                    "3. Verify the error message in the response."
                ),
                "data": (
                    'POST /api/auth/register\n'
                    '{\n'
                    '  "firstName": "Jane",\n'
                    '  "lastName": "Smith",\n'
                    '  "email": "john.doe@example.com",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 400 Bad Request.\n"
                    "Response body: { success: false, message: 'Email already in use' } (or equivalent)."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-003",
                "title": "Registration with Invalid Email Format",
                "module": "User Authentication",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send POST /api/auth/register with a malformed email address.\n"
                    "2. Verify the HTTP response code.\n"
                    "3. Verify validation errors are returned."
                ),
                "data": (
                    'POST /api/auth/register\n'
                    '{\n'
                    '  "firstName": "Jane",\n'
                    '  "lastName": "Smith",\n'
                    '  "email": "not-an-email",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Response body: { success: false, error: [ { msg: 'Please provide a valid email', ... } ] }."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-004",
                "title": "Registration with Password Shorter Than 6 Characters",
                "module": "User Authentication",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send POST /api/auth/register with a 5-character password.\n"
                    "2. Verify the HTTP response code.\n"
                    "3. Verify the validation error references password length."
                ),
                "data": (
                    'POST /api/auth/register\n'
                    '{\n'
                    '  "firstName": "Jane",\n'
                    '  "lastName": "Smith",\n'
                    '  "email": "jane.smith@example.com",\n'
                    '  "password": "Pass1"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Response body contains validation error: 'Password must be at least 6 characters' (or equivalent)."
                ),
                "priority": "Medium",
                "type": "Boundary",
            },
            {
                "id": "TC-AUTH-005",
                "title": "Registration with Missing Required Fields",
                "module": "User Authentication",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send POST /api/auth/register omitting the firstName field.\n"
                    "2. Verify the HTTP response code.\n"
                    "3. Verify the validation error identifies the missing field."
                ),
                "data": (
                    'POST /api/auth/register\n'
                    '{\n'
                    '  "lastName": "Smith",\n'
                    '  "email": "jane.smith@example.com",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Response contains validation error indicating firstName is required."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-006",
                "title": "Successful User Login",
                "module": "User Authentication",
                "preconditions": "A user with email 'john.doe@example.com' and password 'Password123' exists.",
                "steps": (
                    "1. Send POST /api/auth/login with correct credentials.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify a JWT token is returned.\n"
                    "4. Confirm password is not included in the response."
                ),
                "data": (
                    'POST /api/auth/login\n'
                    '{\n'
                    '  "email": "john.doe@example.com",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response body: { success: true, data: { user: {...}, token: '...' } }.\n"
                    "data.token is a non-empty JWT string.\n"
                    "data.user does NOT contain password."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-AUTH-007",
                "title": "Login with Incorrect Password",
                "module": "User Authentication",
                "preconditions": "User 'john.doe@example.com' exists.",
                "steps": (
                    "1. Send POST /api/auth/login with the correct email but wrong password.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the error message."
                ),
                "data": (
                    'POST /api/auth/login\n'
                    '{\n'
                    '  "email": "john.doe@example.com",\n'
                    '  "password": "WrongPassword"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 401 Unauthorized.\n"
                    "Response body: { success: false, message: 'Invalid credentials' } (or equivalent)."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-008",
                "title": "Login with Non-Existent Email",
                "module": "User Authentication",
                "preconditions": "No user with email 'ghost@example.com' exists.",
                "steps": (
                    "1. Send POST /api/auth/login with a non-registered email.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the error message."
                ),
                "data": (
                    'POST /api/auth/login\n'
                    '{\n'
                    '  "email": "ghost@example.com",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 401 Unauthorized.\n"
                    "Response body: { success: false, message: 'Invalid credentials' }."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-009",
                "title": "Login with Invalid Email Format",
                "module": "User Authentication",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send POST /api/auth/login with a malformed email.\n"
                    "2. Verify HTTP response code."
                ),
                "data": (
                    'POST /api/auth/login\n'
                    '{\n'
                    '  "email": "notanemail",\n'
                    '  "password": "Password123"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Response contains validation error for email field."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-010",
                "title": "Get User Profile with Valid Token",
                "module": "User Authentication",
                "preconditions": "User is registered and a valid JWT token has been obtained via login.",
                "steps": (
                    "1. Send GET /api/auth/profile with Authorization: Bearer <valid_token>.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the returned user data matches the logged-in user."
                ),
                "data": (
                    'GET /api/auth/profile\n'
                    'Header: Authorization: Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response body: { success: true, data: { user: { firstName, lastName, email, ... } } }.\n"
                    "Returned email matches the authenticated user's email."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-AUTH-011",
                "title": "Get User Profile Without Token",
                "module": "User Authentication",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send GET /api/auth/profile without an Authorization header.\n"
                    "2. Verify HTTP response code."
                ),
                "data": "GET /api/auth/profile\n(No Authorization header)",
                "expected": (
                    "HTTP 401 Unauthorized.\n"
                    "Response body: { success: false, message: 'No token provided' or equivalent }."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-AUTH-012",
                "title": "Get User Profile with Expired/Invalid Token",
                "module": "User Authentication",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send GET /api/auth/profile with an expired or tampered JWT.\n"
                    "2. Verify HTTP response code."
                ),
                "data": (
                    'GET /api/auth/profile\n'
                    'Header: Authorization: Bearer invalidtoken123'
                ),
                "expected": (
                    "HTTP 401 Unauthorized.\n"
                    "Response body: { success: false, message: 'Invalid token' or equivalent }."
                ),
                "priority": "High",
                "type": "Negative",
            },
        ],
    },
    {
        "id": "MODULE 2",
        "name": "A/L Profile Management",
        "cases": [
            {
                "id": "TC-PROF-001",
                "title": "Create A/L Profile with All Required Fields",
                "module": "A/L Profile Management",
                "preconditions": "User is authenticated. Valid JWT token available.",
                "steps": (
                    "1. Send POST /api/profiles with all required and optional fields.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the created profile is returned."
                ),
                "data": (
                    'POST /api/profiles  (Bearer token required)\n'
                    '{\n'
                    '  "age": "22",\n'
                    '  "gender": "Male",\n'
                    '  "nativeLanguage": "Sinhala",\n'
                    '  "preferredLanguage": "English",\n'
                    '  "olResults": "9A passes",\n'
                    '  "alStream": "Physical Science",\n'
                    '  "interestArea": "Information Technology",\n'
                    '  "careerGoal": "Software Engineer",\n'
                    '  "monthlyIncome": "50000",\n'
                    '  "fundingMethod": "Self-funded",\n'
                    '  "availability": "Weekday",\n'
                    '  "completionPeriod": "3-4 years",\n'
                    '  "studyMethod": "Hybrid",\n'
                    '  "currentLocation": "Colombo",\n'
                    '  "preferredLocations": "Colombo, Kandy"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 201 Created.\n"
                    "Response body: { success: true, data: { profile: { _id, userId, age, gender, ... } } }.\n"
                    "profile._id is a valid MongoDB ObjectId.\n"
                    "profile.userId matches the authenticated user's ID."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-PROF-002",
                "title": "Create Profile Without Authentication",
                "module": "A/L Profile Management",
                "preconditions": "Application server is running. No auth token provided.",
                "steps": (
                    "1. Send POST /api/profiles without an Authorization header.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST /api/profiles\n(No Authorization header)\n{ "age": "22", "gender": "Male", ... }',
                "expected": (
                    "HTTP 401 Unauthorized.\n"
                    "Response body: { success: false, message: 'No token provided' }."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-PROF-003",
                "title": "Create Profile with Invalid Gender Enum Value",
                "module": "A/L Profile Management",
                "preconditions": "User is authenticated.",
                "steps": (
                    "1. Send POST /api/profiles with gender set to an invalid value.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify validation error is returned."
                ),
                "data": (
                    'POST /api/profiles  (Bearer token required)\n'
                    '{\n'
                    '  "age": "22",\n'
                    '  "gender": "Unknown",\n'
                    '  "nativeLanguage": "English",\n'
                    '  ...\n'
                    '}'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Response contains validation error: gender must be one of ['Male', 'Female', 'Other']."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-PROF-004",
                "title": "Create Profile with Invalid Interest Area Value",
                "module": "A/L Profile Management",
                "preconditions": "User is authenticated.",
                "steps": (
                    "1. Send POST /api/profiles with an interestArea value not in the allowed enum list.\n"
                    "2. Verify HTTP response code."
                ),
                "data": (
                    'POST /api/profiles  (Bearer token required)\n'
                    '{ ..., "interestArea": "Cooking", ... }'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error indicates interestArea must be one of the 7 allowed values."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-PROF-005",
                "title": "Get All Profiles for Authenticated User",
                "module": "A/L Profile Management",
                "preconditions": "User is authenticated. At least one profile exists for this user.",
                "steps": (
                    "1. Send GET /api/profiles with a valid Bearer token.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify all returned profiles belong to the authenticated user."
                ),
                "data": "GET /api/profiles\nHeader: Authorization: Bearer <valid_token>",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response body: { success: true, data: { profiles: [...] } }.\n"
                    "Every profile in the array has a userId matching the authenticated user."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-PROF-006",
                "title": "Get Profile by Valid ID",
                "module": "A/L Profile Management",
                "preconditions": "Profile with ID '64f1c2d3e4b0a1b2c3d4e5f6' exists.",
                "steps": (
                    "1. Send GET /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify returned profile ID matches the requested ID."
                ),
                "data": "GET /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response body: { success: true, data: { profile: { _id: '64f1c2d3e4b0a1b2c3d4e5f6', ... } } }."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-PROF-007",
                "title": "Get Profile by Non-Existent ID",
                "module": "A/L Profile Management",
                "preconditions": "No profile with ID '64f1c2d3e4b0a1b2c3d4e500' exists.",
                "steps": (
                    "1. Send GET /api/profiles/64f1c2d3e4b0a1b2c3d4e500.\n"
                    "2. Verify HTTP response code."
                ),
                "data": "GET /api/profiles/64f1c2d3e4b0a1b2c3d4e500",
                "expected": "HTTP 404 Not Found.\nResponse body: { success: false, message: 'Profile not found' }.",
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-PROF-008",
                "title": "Get Profile with Invalid MongoDB ObjectId",
                "module": "A/L Profile Management",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send GET /api/profiles/not-an-object-id.\n"
                    "2. Verify HTTP response code."
                ),
                "data": "GET /api/profiles/not-an-object-id",
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Response body contains validation error: id must be a valid MongoDB ObjectId."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-PROF-009",
                "title": "Update Profile with Valid Data",
                "module": "A/L Profile Management",
                "preconditions": "User is authenticated. Profile with ID '64f1c2d3e4b0a1b2c3d4e5f6' belongs to the authenticated user.",
                "steps": (
                    "1. Send PUT /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6 with updated fields.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the updated fields are reflected in the response."
                ),
                "data": (
                    'PUT /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6  (Bearer token)\n'
                    '{ "careerGoal": "Data Scientist", "studyMethod": "Online" }'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response body: { success: true, data: { profile: { careerGoal: 'Data Scientist', studyMethod: 'Online', ... } } }."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-PROF-010",
                "title": "Update Another User's Profile (Forbidden)",
                "module": "A/L Profile Management",
                "preconditions": "User A is authenticated. Profile '64f1c2d3e4b0a1b2c3d4e5f7' belongs to User B.",
                "steps": (
                    "1. Authenticate as User A.\n"
                    "2. Send PUT /api/profiles/64f1c2d3e4b0a1b2c3d4e5f7 with updated data.\n"
                    "3. Verify HTTP response code."
                ),
                "data": 'PUT /api/profiles/64f1c2d3e4b0a1b2c3d4e5f7  (User A Bearer token)\n{ "careerGoal": "Hacker" }',
                "expected": (
                    "HTTP 403 Forbidden or 404 Not Found.\n"
                    "Response body: { success: false, message: 'Access denied' or 'Profile not found' }."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-PROF-011",
                "title": "Delete Own Profile",
                "module": "A/L Profile Management",
                "preconditions": "User is authenticated. Profile '64f1c2d3e4b0a1b2c3d4e5f6' belongs to authenticated user.",
                "steps": (
                    "1. Send DELETE /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6 with valid token.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Send GET /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6 to confirm deletion."
                ),
                "data": "DELETE /api/profiles/64f1c2d3e4b0a1b2c3d4e5f6  (Bearer token)",
                "expected": (
                    "HTTP 200 OK for DELETE.\n"
                    "Subsequent GET returns HTTP 404."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-PROF-012",
                "title": "Delete Another User's Profile (Forbidden)",
                "module": "A/L Profile Management",
                "preconditions": "User A authenticated. Profile '64f1c2d3e4b0a1b2c3d4e5f7' belongs to User B.",
                "steps": (
                    "1. Authenticate as User A.\n"
                    "2. Send DELETE /api/profiles/64f1c2d3e4b0a1b2c3d4e5f7."
                ),
                "data": "DELETE /api/profiles/64f1c2d3e4b0a1b2c3d4e5f7  (User A Bearer token)",
                "expected": "HTTP 403 Forbidden or 404 Not Found.\nProfile record is NOT deleted from the database.",
                "priority": "High",
                "type": "Negative",
            },
        ],
    },
    {
        "id": "MODULE 3",
        "name": "Course Recommendation",
        "cases": [
            {
                "id": "TC-REC-001",
                "title": "Get Recommendations with Complete User Profile",
                "module": "Course Recommendation",
                "preconditions": "Python FastAPI server is running on port 8000. ChromaDB vector store is populated.",
                "steps": (
                    "1. Send POST /recommend with a fully populated UserProfile payload.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the recommendations array contains entries."
                ),
                "data": (
                    'POST http://localhost:8000/recommend\n'
                    '{\n'
                    '  "age": "22",\n'
                    '  "native_language": "Sinhala",\n'
                    '  "preferred_language": "English",\n'
                    '  "ol_results": "9A passes",\n'
                    '  "al_stream": "Physical Science",\n'
                    '  "interest_area": "Information Technology",\n'
                    '  "career_goal": "Software Engineer",\n'
                    '  "income": "50000",\n'
                    '  "study_method": "Hybrid",\n'
                    '  "availability": "Weekday",\n'
                    '  "completion_period": "3-4 years",\n'
                    '  "current_location": "Colombo",\n'
                    '  "preferred_locations": "Colombo"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { status: 'success', recommendations: [...], warnings: [], errors: [] }.\n"
                    "recommendations array has at most 10 entries, each with: id, rank, course_name, university, match_score, explanation."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-REC-002",
                "title": "Get Recommendations with Partial Profile",
                "module": "Course Recommendation",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send POST /recommend with only a few fields populated (all fields are optional).\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify recommendations are still returned (possibly with warnings)."
                ),
                "data": (
                    'POST http://localhost:8000/recommend\n'
                    '{\n'
                    '  "career_goal": "Doctor",\n'
                    '  "al_stream": "Bio Science"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "recommendations array is non-empty.\n"
                    "warnings array may contain messages about missing profile fields."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-REC-003",
                "title": "Get Recommendations with Empty Profile",
                "module": "Course Recommendation",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send POST /recommend with an empty JSON object.\n"
                    "2. Verify the system responds without crashing."
                ),
                "data": "POST http://localhost:8000/recommend\n{}",
                "expected": (
                    "HTTP 200 OK (or 422 if validation requires at least one field).\n"
                    "No server error (5xx). Either returns generic recommendations or a meaningful error message."
                ),
                "priority": "Medium",
                "type": "Boundary",
            },
            {
                "id": "TC-REC-004",
                "title": "Validate Recommendation Response Structure",
                "module": "Course Recommendation",
                "preconditions": "Python FastAPI server is running. Valid profile data available.",
                "steps": (
                    "1. Send POST /recommend with a complete profile.\n"
                    "2. For each item in recommendations, verify all required fields are present."
                ),
                "data": "See TC-REC-001 test data.",
                "expected": (
                    "Each recommendation object contains: id (UUID string), rank (integer), course_name (string), "
                    "university (string), location (string), match_score (float 0-1), explanation (string), tags (array)."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-REC-005",
                "title": "Recommendations Return At Most 10 Results",
                "module": "Course Recommendation",
                "preconditions": "Python FastAPI server is running. Vector store contains more than 10 courses.",
                "steps": (
                    "1. Send POST /recommend with a complete profile.\n"
                    "2. Count the items in the recommendations array."
                ),
                "data": "See TC-REC-001 test data.",
                "expected": (
                    "recommendations.length <= 10.\n"
                    "Results are ordered by rank (rank 1 = highest match_score)."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-REC-006",
                "title": "Recommendations Respect Preferred Language Filter",
                "module": "Course Recommendation",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send POST /recommend with preferred_language set to 'Sinhala'.\n"
                    "2. Inspect the study_language field of returned recommendations."
                ),
                "data": 'POST http://localhost:8000/recommend\n{ "preferred_language": "Sinhala", "career_goal": "Teacher" }',
                "expected": (
                    "HTTP 200 OK.\n"
                    "Recommendations with Sinhala as study_language are ranked higher or filtered accordingly."
                ),
                "priority": "Low",
                "type": "Positive",
            },
        ],
    },
    {
        "id": "MODULE 4",
        "name": "A/L Stream Prediction",
        "cases": [
            {
                "id": "TC-STRM-001",
                "title": "Get Available A/L Streams List",
                "module": "A/L Stream Prediction",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send GET /api/predict/streams.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the returned list includes all expected streams."
                ),
                "data": "GET http://localhost:8000/api/predict/streams",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response includes streams: 'Bio Science', 'Physical Science', 'Commerce', 'Arts', "
                    "'Engineering Technology', 'Bio-systems Technology'."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-STRM-002",
                "title": "Get Subject Basket Configuration",
                "module": "A/L Stream Prediction",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send GET /api/predict/baskets.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify each stream has a corresponding subject list."
                ),
                "data": "GET http://localhost:8000/api/predict/baskets",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response contains a mapping of stream names to their subject baskets."
                ),
                "priority": "Low",
                "type": "Positive",
            },
            {
                "id": "TC-STRM-003",
                "title": "Predict A/L Stream with Valid O/L Marks",
                "module": "A/L Stream Prediction",
                "preconditions": "Python FastAPI server running. ML model loaded.",
                "steps": (
                    "1. Send POST /api/predict/stream with valid O/L subject marks for all 22 subjects.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the response contains recommended streams with confidence scores."
                ),
                "data": (
                    'POST http://localhost:8000/api/predict/stream\n'
                    '{\n'
                    '  "Sinhala": "A", "English": "A", "Maths": "A",\n'
                    '  "Science": "A", "History": "B", "Religion": "A",\n'
                    '  "ICT": "A", ...\n'
                    '  (all 22 O/L subjects with grade A-F)\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response contains: recommended_streams (list), each with stream name and confidence score.\n"
                    "At least one stream is recommended."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-STRM-004",
                "title": "Predict Stream with Borderline Grades",
                "module": "A/L Stream Prediction",
                "preconditions": "Python FastAPI server running.",
                "steps": (
                    "1. Send POST /api/predict/stream where subject grades are all 'C' (borderline passing).\n"
                    "2. Verify system returns a prediction without error."
                ),
                "data": 'POST http://localhost:8000/api/predict/stream\n{ All 22 subjects set to grade "C" }',
                "expected": (
                    "HTTP 200 OK.\n"
                    "A stream recommendation is still returned, possibly Arts or a less competitive stream."
                ),
                "priority": "Medium",
                "type": "Boundary",
            },
            {
                "id": "TC-STRM-005",
                "title": "Predict Stream with Missing Subject Grades",
                "module": "A/L Stream Prediction",
                "preconditions": "Python FastAPI server running.",
                "steps": (
                    "1. Send POST /api/predict/stream with only 10 out of 22 subjects provided.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST http://localhost:8000/api/predict/stream\n{ Only 10 subjects provided }',
                "expected": (
                    "HTTP 422 Unprocessable Entity or HTTP 200 with a warning.\n"
                    "No server 500 error."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-STRM-006",
                "title": "Predict Stream with Invalid Grade Letter",
                "module": "A/L Stream Prediction",
                "preconditions": "Python FastAPI server running.",
                "steps": (
                    "1. Send POST /api/predict/stream with an invalid grade letter (e.g., 'Z') for one subject.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST http://localhost:8000/api/predict/stream\n{ ..., "Maths": "Z", ... }',
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error indicates grade must be one of A, B, C, D, E, F."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
        ],
    },
    {
        "id": "MODULE 5",
        "name": "Aptitude Quiz",
        "cases": [
            {
                "id": "TC-QUIZ-001",
                "title": "Get Available Quiz Streams",
                "module": "Aptitude Quiz",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send GET /api/quiz/streams.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the response is a non-empty list of stream names."
                ),
                "data": "GET http://localhost:8000/api/quiz/streams",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response is a list containing at least one stream name string."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-QUIZ-002",
                "title": "Get Stream Information for Valid Stream",
                "module": "Aptitude Quiz",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send GET /api/quiz/stream-info/Physical%20Science.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify stream info is returned."
                ),
                "data": "GET http://localhost:8000/api/quiz/stream-info/Physical%20Science",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response contains stream name and descriptive information."
                ),
                "priority": "Low",
                "type": "Positive",
            },
            {
                "id": "TC-QUIZ-003",
                "title": "Get Stream Info for Invalid Stream Name",
                "module": "Aptitude Quiz",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send GET /api/quiz/stream-info/UnknownStream.\n"
                    "2. Verify HTTP response code."
                ),
                "data": "GET http://localhost:8000/api/quiz/stream-info/UnknownStream",
                "expected": (
                    "HTTP 404 Not Found.\n"
                    "Response body: { detail: 'Stream not found' } or equivalent."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-QUIZ-004",
                "title": "Generate Quiz for Valid Stream",
                "module": "Aptitude Quiz",
                "preconditions": "Python FastAPI server is running. LLM/quiz generator is available.",
                "steps": (
                    "1. Send POST /api/quiz/generate with a valid stream and num_questions=5.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify exactly 5 questions are returned.\n"
                    "4. Verify each question has id, question, and options fields."
                ),
                "data": (
                    'POST http://localhost:8000/api/quiz/generate\n'
                    '{ "stream": "Physical Science", "num_questions": 5 }'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { stream: 'Physical Science', questions: [ (5 items) ], answers: [ (5 items) ] }.\n"
                    "Each question has: id, question (string), options (array of choices)."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-QUIZ-005",
                "title": "Generate Quiz with Out-of-Range Question Count",
                "module": "Aptitude Quiz",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send POST /api/quiz/generate with num_questions=25 (exceeds max 20).\n"
                    "2. Verify HTTP response code."
                ),
                "data": (
                    'POST http://localhost:8000/api/quiz/generate\n'
                    '{ "stream": "Physical Science", "num_questions": 25 }'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error: num_questions must be between 1 and 20."
                ),
                "priority": "Medium",
                "type": "Boundary",
            },
            {
                "id": "TC-QUIZ-006",
                "title": "Submit Quiz with All Correct Answers",
                "module": "Aptitude Quiz",
                "preconditions": "Quiz was generated and correct answers are known.",
                "steps": (
                    "1. Generate a 5-question quiz and note the correct_answers.\n"
                    "2. Send POST /api/quiz/submit with answers matching correct_answers.\n"
                    "3. Verify score equals 5 and percentage equals 100."
                ),
                "data": (
                    'POST http://localhost:8000/api/quiz/submit\n'
                    '{\n'
                    '  "answers": ["A", "B", "C", "D", "A"],\n'
                    '  "correct_answers": ["A", "B", "C", "D", "A"]\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { score: 5, total: 5, percentage: 100.0, results: [ all is_correct: true ] }."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-QUIZ-007",
                "title": "Submit Quiz with All Wrong Answers",
                "module": "Aptitude Quiz",
                "preconditions": "Quiz was generated and correct answers are known.",
                "steps": (
                    "1. Generate a 5-question quiz and note correct_answers.\n"
                    "2. Send POST /api/quiz/submit with intentionally wrong answers.\n"
                    "3. Verify score equals 0 and percentage equals 0."
                ),
                "data": (
                    'POST http://localhost:8000/api/quiz/submit\n'
                    '{\n'
                    '  "answers": ["B", "A", "D", "C", "B"],\n'
                    '  "correct_answers": ["A", "B", "C", "D", "A"]\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { score: 0, total: 5, percentage: 0.0, results: [ all is_correct: false ] }."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-QUIZ-008",
                "title": "Submit Quiz with Mismatched Answer Count",
                "module": "Aptitude Quiz",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send POST /api/quiz/submit where answers has 3 items but correct_answers has 5 items.\n"
                    "2. Verify the system handles the mismatch gracefully."
                ),
                "data": (
                    'POST http://localhost:8000/api/quiz/submit\n'
                    '{\n'
                    '  "answers": ["A", "B", "C"],\n'
                    '  "correct_answers": ["A", "B", "C", "D", "A"]\n'
                    '}'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity or HTTP 400 Bad Request.\n"
                    "Response indicates answers and correct_answers must have the same length."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
        ],
    },
    {
        "id": "MODULE 6",
        "name": "University Prediction",
        "cases": [
            {
                "id": "TC-UPRED-001",
                "title": "Predict Universities with Valid Student Profile",
                "module": "University Prediction",
                "preconditions": "Python FastAPI server running. ML models loaded.",
                "steps": (
                    "1. Send POST /predict with a fully populated student profile.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify top_predictions contains up to 10 universities."
                ),
                "data": (
                    'POST http://localhost:8000/predict\n'
                    '{\n'
                    '  "Year": 2024.0,\n'
                    '  "Stream": "Physical Science",\n'
                    '  "Subject_1": "Physics", "Grade_1": "A",\n'
                    '  "Subject_2": "Combined Maths", "Grade_2": "A",\n'
                    '  "Subject_3": "Chemistry", "Grade_3": "B",\n'
                    '  "Z_Score": 1.8,\n'
                    '  "Island_Rank": 1200,\n'
                    '  "District": "Colombo",\n'
                    '  "Gen_Test": 4.5,\n'
                    '  "Sinhala/Tamil": "B", "English": "A", "Maths": "A", "Science": "A",\n'
                    '  "q1_science_tech": 5, "q2_problem_solving": 4, ...(all 12 q fields),\n'
                    '  "Course": "Computer Science"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { top_predictions: [ {university, probability}, ... ], course: 'Computer Science', input_summary: {...} }.\n"
                    "top_predictions contains up to 10 entries sorted by probability descending."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UPRED-002",
                "title": "Predict Universities with Missing Required Field",
                "module": "University Prediction",
                "preconditions": "Python FastAPI server is running.",
                "steps": (
                    "1. Send POST /predict omitting the 'Course' field.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST http://localhost:8000/predict\n{ (all fields except Course) }',
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error indicates Course is required."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-UPRED-003",
                "title": "Batch Predict Universities for Multiple Students",
                "module": "University Prediction",
                "preconditions": "Python FastAPI server running.",
                "steps": (
                    "1. Send POST /predict/batch with an array of 2 student profiles.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify two result sets are returned."
                ),
                "data": (
                    'POST http://localhost:8000/predict/batch\n'
                    '[ { student_1_profile }, { student_2_profile } ]'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response is an array of 2 prediction result objects."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-UPRED-004",
                "title": "Combined Course and University Prediction",
                "module": "University Prediction",
                "preconditions": "Python FastAPI server running. Cutoff data available.",
                "steps": (
                    "1. Send POST /predict/combined with a student profile and top_n_courses=3.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify 3 courses are returned each with university predictions."
                ),
                "data": (
                    'POST http://localhost:8000/predict/combined\n'
                    '{ ...student profile fields..., "top_n_courses": 3 }'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { results: [ 3 items ], input_summary: {...} }.\n"
                    "Each item has: rank, course, course_score, university, top_university_predictions."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UPRED-005",
                "title": "Combined Prediction with Boundary top_n_courses Values",
                "module": "University Prediction",
                "preconditions": "Python FastAPI server running.",
                "steps": (
                    "1. Send POST /predict/combined with top_n_courses=1 (minimum).\n"
                    "2. Verify response contains exactly 1 result.\n"
                    "3. Send POST /predict/combined with top_n_courses=20 (maximum).\n"
                    "4. Verify response contains up to 20 results."
                ),
                "data": (
                    'POST http://localhost:8000/predict/combined\n'
                    '{ ...student profile..., "top_n_courses": 1 }  (then repeat with 20)'
                ),
                "expected": (
                    "Both requests: HTTP 200 OK.\n"
                    "top_n_courses=1 → results.length == 1.\n"
                    "top_n_courses=20 → results.length <= 20."
                ),
                "priority": "Medium",
                "type": "Boundary",
            },
            {
                "id": "TC-UPRED-006",
                "title": "Predict with Out-of-Range Z-Score",
                "module": "University Prediction",
                "preconditions": "Python FastAPI server running.",
                "steps": (
                    "1. Send POST /predict with Z_Score set to -5.0 (invalid negative value).\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST http://localhost:8000/predict\n{ ...all fields..., "Z_Score": -5.0 }',
                "expected": (
                    "HTTP 422 Unprocessable Entity or HTTP 200 with an error flag.\n"
                    "No 500 server error."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
        ],
    },
    {
        "id": "MODULE 7",
        "name": "Roadmap Generation",
        "cases": [
            {
                "id": "TC-ROAD-001",
                "title": "Generate Roadmap with Valid Profile and Course",
                "module": "Roadmap Generation",
                "preconditions": (
                    "User is authenticated. A valid AL profile exists (profileId: '64f1c2d3e4b0a1b2c3d4e5f6'). "
                    "Python backend and LLM API are reachable."
                ),
                "steps": (
                    "1. Send POST /api/diagram with profileId and courseId.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the generated roadmap contains exactly 6 stages."
                ),
                "data": (
                    'POST /api/diagram  (Bearer token)\n'
                    '{\n'
                    '  "profileId": "64f1c2d3e4b0a1b2c3d4e5f6",\n'
                    '  "courseId": "cs-university-of-colombo"\n'
                    '}'
                ),
                "expected": (
                    "HTTP 201 Created.\n"
                    "Response: { success: true, data: { roadmap: { _id, userId, profileId, courseId, roadmap: [...6 stages], status: 'success' } } }.\n"
                    "roadmap array has exactly 6 items, each with: id, title, goal, duration, description, actionPlan, resources, successCriteria."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-ROAD-002",
                "title": "Generate Roadmap Without Authentication",
                "module": "Roadmap Generation",
                "preconditions": "Application server is running.",
                "steps": (
                    "1. Send POST /api/diagram without an Authorization header.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST /api/diagram\n(No Authorization header)\n{ "profileId": "...", "courseId": "..." }',
                "expected": "HTTP 401 Unauthorized.\nResponse: { success: false, message: 'No token provided' }.",
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-ROAD-003",
                "title": "Generate Roadmap with Invalid ProfileId Format",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated.",
                "steps": (
                    "1. Send POST /api/diagram with profileId set to a non-ObjectId string.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST /api/diagram  (Bearer token)\n{ "profileId": "not-valid-id", "courseId": "cs-uoc" }',
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error: profileId must be a valid MongoDB ObjectId."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-ROAD-004",
                "title": "Generate Roadmap with Non-Existent Profile",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated. Profile ID does not exist in the database.",
                "steps": (
                    "1. Send POST /api/diagram with a valid ObjectId format but non-existent profileId.\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'POST /api/diagram  (Bearer token)\n{ "profileId": "64f1c2d3e4b0a1b2c3000000", "courseId": "cs-uoc" }',
                "expected": (
                    "HTTP 404 Not Found.\n"
                    "Response: { success: false, message: 'Profile not found' }."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-ROAD-005",
                "title": "Get All Roadmaps for Authenticated User",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated. At least one roadmap exists for this user.",
                "steps": (
                    "1. Send GET /api/diagram with valid Bearer token.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify all returned roadmaps belong to the authenticated user."
                ),
                "data": "GET /api/diagram\nHeader: Authorization: Bearer <valid_token>",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { success: true, data: { roadmaps: [...] } }.\n"
                    "Every roadmap has userId matching the authenticated user."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-ROAD-006",
                "title": "Get Roadmap by Valid ID",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated. Roadmap with given ID exists.",
                "steps": (
                    "1. Send GET /api/diagram/64f1c2d3e4b0a1b2c3d4e5f9 with Bearer token.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the returned roadmap ID matches."
                ),
                "data": "GET /api/diagram/64f1c2d3e4b0a1b2c3d4e5f9  (Bearer token)",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { success: true, data: { roadmap: { _id: '64f1c2d3e4b0a1b2c3d4e5f9', ... } } }."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-ROAD-007",
                "title": "Get Roadmap by Profile and Course",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated. Roadmap for this profile-course combination exists.",
                "steps": (
                    "1. Send GET /api/diagram/profile/64f1c2d3e4b0a1b2c3d4e5f6/course/cs-uoc.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the returned roadmap matches the profile and course."
                ),
                "data": "GET /api/diagram/profile/64f1c2d3e4b0a1b2c3d4e5f6/course/cs-uoc  (Bearer token)",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Returned roadmap has profileId='64f1c2d3e4b0a1b2c3d4e5f6' and courseId='cs-uoc'."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-ROAD-008",
                "title": "Get Roadmap by Non-Existent ID",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated. No roadmap with given ID exists.",
                "steps": (
                    "1. Send GET /api/diagram/64f1c2d3e4b0a1b2c3000000.\n"
                    "2. Verify HTTP response code."
                ),
                "data": "GET /api/diagram/64f1c2d3e4b0a1b2c3000000  (Bearer token)",
                "expected": "HTTP 404 Not Found.\nResponse: { success: false, message: 'Roadmap not found' }.",
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-ROAD-009",
                "title": "Delete Roadmap",
                "module": "Roadmap Generation",
                "preconditions": "User is authenticated. Roadmap '64f1c2d3e4b0a1b2c3d4e5f9' belongs to the user.",
                "steps": (
                    "1. Send DELETE /api/diagram/64f1c2d3e4b0a1b2c3d4e5f9 with Bearer token.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Send GET /api/diagram/64f1c2d3e4b0a1b2c3d4e5f9 and verify it is no longer accessible."
                ),
                "data": "DELETE /api/diagram/64f1c2d3e4b0a1b2c3d4e5f9  (Bearer token)",
                "expected": (
                    "HTTP 200 OK for DELETE.\n"
                    "Subsequent GET returns HTTP 404.\n"
                    "Record is soft-deleted (isDeleted: true in database)."
                ),
                "priority": "High",
                "type": "Positive",
            },
        ],
    },
    {
        "id": "MODULE 8",
        "name": "Milestone Tracking",
        "cases": [
            {
                "id": "TC-MILE-001",
                "title": "Create Milestone with Valid 6 Stages",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. A roadmap with roadmapId '64f1c2d3e4b0a1b2c3d4e5f9' exists.",
                "steps": (
                    "1. Send POST /api/milestones with all required fields including exactly 6 stages.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the created milestone is returned with progressPercentage of 0."
                ),
                "data": (
                    'POST /api/milestones  (Bearer token)\n'
                    '{\n'
                    '  "profileId": "64f1c2d3e4b0a1b2c3d4e5f6",\n'
                    '  "roadmapId": "64f1c2d3e4b0a1b2c3d4e5f9",\n'
                    '  "courseId": "cs-uoc",\n'
                    '  "courseName": "Computer Science",\n'
                    '  "university": "University of Colombo",\n'
                    '  "careerGoal": "Software Engineer",\n'
                    '  "stages": [\n'
                    '    { "stepId": 1, "stepTitle": "Foundation", "stepGoal": "...", "description": "...", "duration": "3 months" },\n'
                    '    { "stepId": 2, "stepTitle": "Core Skills", ... },\n'
                    '    ... (6 stages total)\n'
                    '  ]\n'
                    '}'
                ),
                "expected": (
                    "HTTP 201 Created.\n"
                    "Response: { success: true, data: { milestone: { _id, progressPercentage: 0, overallStatus: 'pending', stages: [6 items] } } }."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-002",
                "title": "Create Milestone with Wrong Stage Count",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated.",
                "steps": (
                    "1. Send POST /api/milestones with only 4 stages in the stages array.\n"
                    "2. Verify HTTP response code."
                ),
                "data": (
                    'POST /api/milestones  (Bearer token)\n'
                    '{ ...required fields..., "stages": [ stage1, stage2, stage3, stage4 ] }'
                ),
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error: stages must contain exactly 6 items."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
            {
                "id": "TC-MILE-003",
                "title": "Get All Milestones for Authenticated User",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. At least one milestone exists.",
                "steps": (
                    "1. Send GET /api/milestones with Bearer token.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify all milestones belong to the authenticated user."
                ),
                "data": "GET /api/milestones  (Bearer token)",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { success: true, data: { milestones: [...] } }.\n"
                    "Every milestone has userId matching the authenticated user."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-004",
                "title": "Get Milestone by Roadmap ID",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Milestone linked to roadmap '64f1c2d3e4b0a1b2c3d4e5f9' exists.",
                "steps": (
                    "1. Send GET /api/milestones/roadmap/64f1c2d3e4b0a1b2c3d4e5f9.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the returned milestone references the correct roadmapId."
                ),
                "data": "GET /api/milestones/roadmap/64f1c2d3e4b0a1b2c3d4e5f9  (Bearer token)",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Returned milestone has roadmapId = '64f1c2d3e4b0a1b2c3d4e5f9'."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-005",
                "title": "Get Milestone Progress Metrics",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Milestone for roadmap '64f1c2d3e4b0a1b2c3d4e5f9' exists with 2 completed stages.",
                "steps": (
                    "1. Send GET /api/milestones/progress/64f1c2d3e4b0a1b2c3d4e5f9.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify progress metrics match the actual completion state."
                ),
                "data": "GET /api/milestones/progress/64f1c2d3e4b0a1b2c3d4e5f9  (Bearer token)",
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response: { total: 6, completed: 2, inProgress: 0, notStarted: 4, percentage: 33.33 } (or similar)."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-006",
                "title": "Start a Milestone Stage",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Milestone '64f1c2d3e4b0a1b2c3d4e5fa' exists with stage 1 in 'pending' status.",
                "steps": (
                    "1. Send PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/start with stepId=1.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify stage 1 status changes to 'in_progress'."
                ),
                "data": 'PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/start  (Bearer token)\n{ "stepId": 1 }',
                "expected": (
                    "HTTP 200 OK.\n"
                    "Stage 1 has status: 'in_progress' and startedAt is set to current timestamp."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-007",
                "title": "Start Stage with Out-of-Range stepId",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Milestone exists.",
                "steps": (
                    "1. Send PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/start with stepId=7 (invalid, max is 6).\n"
                    "2. Verify HTTP response code."
                ),
                "data": 'PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/start  (Bearer token)\n{ "stepId": 7 }',
                "expected": (
                    "HTTP 422 Unprocessable Entity.\n"
                    "Validation error: stepId must be an integer between 1 and 6."
                ),
                "priority": "Medium",
                "type": "Boundary",
            },
            {
                "id": "TC-MILE-008",
                "title": "Complete a Stage with Feedback",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Stage 1 of milestone '64f1c2d3e4b0a1b2c3d4e5fa' is 'in_progress'.",
                "steps": (
                    "1. Send PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/complete with stepId=1 and feedback.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify stage 1 status changes to 'completed' and progressPercentage increases."
                ),
                "data": (
                    'PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/complete  (Bearer token)\n'
                    '{ "stepId": 1, "feedback": "Completed all foundation modules successfully." }'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Stage 1 status: 'completed', completedAt is set, feedback is saved.\n"
                    "progressPercentage updates from 0 to ~16.67 (1 of 6 stages)."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-009",
                "title": "Complete Stage Without Providing Feedback",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Stage 1 is 'in_progress'.",
                "steps": (
                    "1. Send PUT /api/milestones/.../stage/complete with stepId=1 but omit the feedback field.\n"
                    "2. Verify the stage is still marked complete (feedback is optional)."
                ),
                "data": 'PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/complete  (Bearer token)\n{ "stepId": 1 }',
                "expected": (
                    "HTTP 200 OK.\n"
                    "Stage 1 status: 'completed'. feedback field is null or empty string."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-010",
                "title": "Update Stage Feedback",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Stage 1 is already completed with existing feedback.",
                "steps": (
                    "1. Send PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/feedback with stepId=1 and new feedback.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify the feedback text is updated."
                ),
                "data": (
                    'PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa/stage/feedback  (Bearer token)\n'
                    '{ "stepId": 1, "feedback": "Updated reflection: challenging but rewarding." }'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Stage 1 feedback now reads 'Updated reflection: challenging but rewarding.'."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-011",
                "title": "Update Overall Milestone Status",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Milestone exists with overallStatus 'pending'.",
                "steps": (
                    "1. Send PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa with overallStatus='in_progress'.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Verify overallStatus is updated."
                ),
                "data": (
                    'PUT /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa  (Bearer token)\n'
                    '{ "overallStatus": "in_progress" }'
                ),
                "expected": (
                    "HTTP 200 OK.\n"
                    "Response milestone has overallStatus: 'in_progress'."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-MILE-012",
                "title": "Soft Delete Milestone",
                "module": "Milestone Tracking",
                "preconditions": "User is authenticated. Milestone '64f1c2d3e4b0a1b2c3d4e5fa' exists.",
                "steps": (
                    "1. Send DELETE /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa.\n"
                    "2. Verify HTTP response code.\n"
                    "3. Send GET /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa to confirm it is no longer returned."
                ),
                "data": "DELETE /api/milestones/64f1c2d3e4b0a1b2c3d4e5fa  (Bearer token)",
                "expected": (
                    "HTTP 200 OK for DELETE.\n"
                    "Subsequent GET returns HTTP 404.\n"
                    "Database record is soft-deleted (isDeleted: true), not physically removed."
                ),
                "priority": "High",
                "type": "Positive",
            },
        ],
    },
    {
        "id": "MODULE 9",
        "name": "Frontend — Authentication UI",
        "cases": [
            {
                "id": "TC-UI-AUTH-001",
                "title": "Login Page Renders Correctly",
                "module": "Frontend — Authentication UI",
                "preconditions": "The Next.js dev server is running. Browser is open.",
                "steps": (
                    "1. Navigate to http://localhost:3000/login.\n"
                    "2. Verify the page loads without console errors.\n"
                    "3. Verify the email field, password field, and Login button are visible."
                ),
                "data": "URL: http://localhost:3000/login",
                "expected": (
                    "Login page renders.\n"
                    "Email and password input fields are present.\n"
                    "Login button is clickable.\n"
                    "No JavaScript console errors."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-AUTH-002",
                "title": "Login Form Inline Validation on Empty Submit",
                "module": "Frontend — Authentication UI",
                "preconditions": "Login page is open.",
                "steps": (
                    "1. Leave email and password fields blank.\n"
                    "2. Click the Login button.\n"
                    "3. Verify inline validation errors appear without submitting to the server."
                ),
                "data": "Email: (empty)\nPassword: (empty)",
                "expected": (
                    "Error messages appear below each field ('Email is required', 'Password is required' or equivalent).\n"
                    "No network request is made to POST /api/auth/login."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-UI-AUTH-003",
                "title": "Successful Login Redirects to Dashboard",
                "module": "Frontend — Authentication UI",
                "preconditions": "A registered user exists with email 'test@example.com' and password 'Password123'.",
                "steps": (
                    "1. Navigate to /login.\n"
                    "2. Enter email: test@example.com and password: Password123.\n"
                    "3. Click the Login button.\n"
                    "4. Verify the browser navigates to the dashboard."
                ),
                "data": "Email: test@example.com\nPassword: Password123",
                "expected": (
                    "Login API returns HTTP 200.\n"
                    "JWT token is stored (cookie or localStorage).\n"
                    "Browser redirects to /dashboard.\n"
                    "Dashboard content is visible."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-AUTH-004",
                "title": "Signup Form Validation",
                "module": "Frontend — Authentication UI",
                "preconditions": "Signup page is open.",
                "steps": (
                    "1. Navigate to /signup.\n"
                    "2. Enter an invalid email format.\n"
                    "3. Enter a password shorter than 6 characters.\n"
                    "4. Click the Sign Up button.\n"
                    "5. Verify inline validation errors appear."
                ),
                "data": "Email: bademail\nPassword: 123\nFirstName: (empty)",
                "expected": (
                    "Validation errors appear for all invalid/missing fields.\n"
                    "No API call is made.\n"
                    "User remains on the signup page."
                ),
                "priority": "High",
                "type": "Negative",
            },
            {
                "id": "TC-UI-AUTH-005",
                "title": "Forgot Password Page Navigation",
                "module": "Frontend — Authentication UI",
                "preconditions": "Login page is open.",
                "steps": (
                    "1. On the login page, click the 'Forgot Password?' link.\n"
                    "2. Verify navigation to the forgot-password page.\n"
                    "3. Verify an email input field is present."
                ),
                "data": "Click 'Forgot Password?' link on /login",
                "expected": (
                    "Browser navigates to /forgot-password (or equivalent route).\n"
                    "An email field and submit button are visible."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
        ],
    },
    {
        "id": "MODULE 10",
        "name": "Frontend — Course Recommendation UI",
        "cases": [
            {
                "id": "TC-UI-REC-001",
                "title": "Course Recommendations List Renders After Profile Submit",
                "module": "Frontend — Course Recommendation UI",
                "preconditions": "User is logged in. The recommendation API is running.",
                "steps": (
                    "1. Navigate to the course recommendation page (e.g., /course-suggestion).\n"
                    "2. Complete the profile form and submit.\n"
                    "3. Verify a list of course recommendations is displayed."
                ),
                "data": "Complete user profile form with all required fields.",
                "expected": (
                    "Loading indicator appears while fetching.\n"
                    "A list of at most 10 recommended courses renders.\n"
                    "Each course card shows course name, university, and match score."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-REC-002",
                "title": "Course Detail Modal Opens and Closes",
                "module": "Frontend — Course Recommendation UI",
                "preconditions": "Course recommendations list is visible.",
                "steps": (
                    "1. Click on a course card in the recommendations list.\n"
                    "2. Verify the course detail modal opens.\n"
                    "3. Click the close button on the modal.\n"
                    "4. Verify the modal closes."
                ),
                "data": "Click on any course recommendation card.",
                "expected": (
                    "Modal opens with course name, university, explanation, career_opportunities, and requirements.\n"
                    "Clicking close dismisses the modal.\n"
                    "The underlying recommendations list is still visible."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-REC-003",
                "title": "Course Selection Updates UI State",
                "module": "Frontend — Course Recommendation UI",
                "preconditions": "Course recommendations list is visible.",
                "steps": (
                    "1. Click 'Select' on a course recommendation card.\n"
                    "2. Verify the card displays a selected state (e.g., highlighted border or checkmark).\n"
                    "3. Verify a 'Continue' or 'Generate Roadmap' button becomes enabled."
                ),
                "data": "Select the top-ranked course recommendation.",
                "expected": (
                    "Selected course card shows a visual indicator (highlighted, checked).\n"
                    "The 'Generate Roadmap' or 'Next' action button activates."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-REC-004",
                "title": "Empty State Message When No Recommendations",
                "module": "Frontend — Course Recommendation UI",
                "preconditions": "Recommendation API returns an empty recommendations array.",
                "steps": (
                    "1. Submit a profile that results in no recommendations (mock/stub the API to return empty array).\n"
                    "2. Verify the UI shows an empty-state message instead of a blank area."
                ),
                "data": "API response: { status: 'success', recommendations: [] }",
                "expected": (
                    "An informative empty-state message is displayed (e.g., 'No courses found for your profile').\n"
                    "No error is thrown. User can edit their profile and retry."
                ),
                "priority": "Medium",
                "type": "Negative",
            },
        ],
    },
    {
        "id": "MODULE 11",
        "name": "Frontend — Roadmap / Diagram UI",
        "cases": [
            {
                "id": "TC-UI-ROAD-001",
                "title": "Visual Roadmap Renders All 6 Stages",
                "module": "Frontend — Roadmap / Diagram UI",
                "preconditions": "User is logged in. A roadmap has been generated and stored.",
                "steps": (
                    "1. Navigate to the roadmap/diagram page (e.g., /diagram).\n"
                    "2. Verify all 6 stage cards are displayed.\n"
                    "3. Verify each stage shows a title, goal, duration, and description."
                ),
                "data": "Navigate to /diagram after roadmap generation.",
                "expected": (
                    "Exactly 6 stage cards are visible.\n"
                    "Each card displays: stage number, title, goal, duration, and description.\n"
                    "No 'undefined' or empty fields visible."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-ROAD-002",
                "title": "Quick Navigation Links Scroll to Correct Stage",
                "module": "Frontend — Roadmap / Diagram UI",
                "preconditions": "Roadmap page is open with 6 stages rendered.",
                "steps": (
                    "1. Locate the quick-navigation panel.\n"
                    "2. Click the link for Stage 4.\n"
                    "3. Verify the page scrolls so that Stage 4 is in view."
                ),
                "data": "Click Stage 4 navigation link.",
                "expected": "The viewport scrolls to Stage 4 card without full page reload.",
                "priority": "Medium",
                "type": "Positive",
            },
            {
                "id": "TC-UI-ROAD-003",
                "title": "Explanation Section Toggles Expand/Collapse",
                "module": "Frontend — Roadmap / Diagram UI",
                "preconditions": "Roadmap page is open.",
                "steps": (
                    "1. Locate a stage that has an explanation section.\n"
                    "2. Click the toggle/expand button.\n"
                    "3. Verify the explanation content expands.\n"
                    "4. Click again to collapse."
                ),
                "data": "Click expand button on Stage 1 explanation section.",
                "expected": (
                    "Expanded: actionPlan, resources, and successCriteria content is visible.\n"
                    "Collapsed: detail content is hidden, only title and goal remain visible."
                ),
                "priority": "Medium",
                "type": "Positive",
            },
        ],
    },
    {
        "id": "MODULE 12",
        "name": "Frontend — Milestone Tracking UI",
        "cases": [
            {
                "id": "TC-UI-MILE-001",
                "title": "Progress Bar Reflects Backend Data",
                "module": "Frontend — Milestone Tracking UI",
                "preconditions": "User is logged in. Milestone exists with 3 of 6 stages completed (50%).",
                "steps": (
                    "1. Navigate to the milestones page (e.g., /milestones).\n"
                    "2. Verify the overall progress bar displays 50%.\n"
                    "3. Verify stage cards that are completed show a 'completed' indicator."
                ),
                "data": "Milestone with progressPercentage: 50.",
                "expected": (
                    "Progress bar visually shows 50%.\n"
                    "3 stage cards have a 'Completed' badge or green indicator.\n"
                    "3 remaining stages show 'Pending' or are grayed out."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-MILE-002",
                "title": "Start Stage Button Triggers API Call",
                "module": "Frontend — Milestone Tracking UI",
                "preconditions": "User is logged in. Stage 1 is in 'pending' status.",
                "steps": (
                    "1. On the milestones page, locate Stage 1 card.\n"
                    "2. Click the 'Start Stage' button.\n"
                    "3. Verify a network request is made to PUT /api/milestones/.../stage/start.\n"
                    "4. Verify Stage 1 card updates to show 'In Progress' status."
                ),
                "data": "Click 'Start Stage' on Stage 1.",
                "expected": (
                    "API call made: PUT /api/milestones/{id}/stage/start  { stepId: 1 }.\n"
                    "Stage 1 card displays 'In Progress' label.\n"
                    "No page reload required."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-MILE-003",
                "title": "Complete Stage Opens Feedback Modal",
                "module": "Frontend — Milestone Tracking UI",
                "preconditions": "User is logged in. Stage 1 is 'in_progress'.",
                "steps": (
                    "1. On the milestones page, locate Stage 1 card in 'In Progress' state.\n"
                    "2. Click the 'Mark Complete' button.\n"
                    "3. Verify a feedback modal appears.\n"
                    "4. Enter feedback text and click 'Confirm'.\n"
                    "5. Verify Stage 1 is updated to 'Completed'."
                ),
                "data": "Click 'Mark Complete' → enter feedback text: 'Done with stage 1'.",
                "expected": (
                    "Feedback modal opens with a text area and Confirm/Cancel buttons.\n"
                    "On Confirm: API call PUT /api/milestones/{id}/stage/complete  { stepId: 1, feedback: '...' }.\n"
                    "Stage 1 card updates to 'Completed' state.\n"
                    "Progress bar increments accordingly."
                ),
                "priority": "High",
                "type": "Positive",
            },
            {
                "id": "TC-UI-MILE-004",
                "title": "Celebration Modal Appears on 100% Completion",
                "module": "Frontend — Milestone Tracking UI",
                "preconditions": "User is logged in. 5 of 6 stages are already completed. Stage 6 is 'in_progress'.",
                "steps": (
                    "1. Click 'Mark Complete' on Stage 6.\n"
                    "2. Enter feedback and confirm.\n"
                    "3. Verify the celebration modal appears."
                ),
                "data": "Complete the final (6th) stage.",
                "expected": (
                    "Progress bar reaches 100%.\n"
                    "A celebration/congratulations modal appears.\n"
                    "overallStatus is updated to 'completed' in the UI."
                ),
                "priority": "High",
                "type": "Positive",
            },
        ],
    },
]


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def set_cell_bg(cell, hex_color: str):
    """Apply background colour to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_tc_table(doc: Document, tc: dict):
    """Render one test case as a 2-column bordered table."""
    FIELDS = [
        ("Test Case ID",    tc["id"]),
        ("Test Case Title", tc["title"]),
        ("Module",          tc["module"]),
        ("Preconditions",   tc["preconditions"]),
        ("Test Steps",      tc["steps"]),
        ("Test Data",       tc["data"]),
        ("Expected Result", tc["expected"]),
        ("Priority",        tc["priority"]),
        ("Test Type",       tc["type"]),
    ]

    table = doc.add_table(rows=len(FIELDS), cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(4.5)
    table.columns[1].width = Cm(13.0)

    HEADER_COLOR = "1F3864"   # dark navy
    HEADER_TXT   = "FFFFFF"
    ALT_COLOR    = "EEF2F7"   # light blue-grey for even rows

    for i, (label, value) in enumerate(FIELDS):
        row = table.rows[i]
        # Label cell
        lc = row.cells[0]
        set_cell_bg(lc, HEADER_COLOR)
        lp = lc.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        lr.font.size = Pt(9)

        # Value cell
        vc = row.cells[1]
        if i % 2 == 1:
            set_cell_bg(vc, ALT_COLOR)
        vp = vc.paragraphs[0]
        vr = vp.add_run(value)
        vr.font.size = Pt(9)

        # Vertical alignment
        lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        vc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_paragraph()   # spacing between test cases


def build_document(output_path: str):
    doc = Document()

    # ---- Page margins ----
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.0)

    # ---- Cover page ----
    doc.add_paragraph()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run("AspireAI — Software Test Case Document")
    tr.bold = True
    tr.font.size = Pt(20)
    tr.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    doc.add_paragraph()
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.add_run("Comprehensive Test Cases for All Major System Modules").font.size = Pt(13)

    doc.add_paragraph()
    meta_lines = [
        f"Project       :  AspireAI Career Guidance Platform",
        f"Document Type :  Software Test Case Document",
        f"Version       :  1.0",
        f"Date          :  {datetime.date.today().strftime('%B %d, %Y')}",
        f"Prepared By   :  QA Team",
    ]
    for line in meta_lines:
        meta_p = doc.add_paragraph(line)
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in meta_p.runs:
            run.font.size = Pt(11)

    doc.add_page_break()

    # ---- Table of contents heading ----
    toc_h = doc.add_heading("Table of Contents", level=1)
    for run in toc_h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    modules_summary = []
    for m in MODULES:
        count = len(m["cases"])
        first_id = m["cases"][0]["id"]
        last_id  = m["cases"][-1]["id"]
        modules_summary.append(
            f"  {m['id']} — {m['name']}  ({first_id} – {last_id}, {count} test cases)"
        )

    total_tcs = sum(len(m["cases"]) for m in MODULES)
    for line in modules_summary:
        p = doc.add_paragraph(line, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(10)

    doc.add_paragraph()
    total_p = doc.add_paragraph(f"Total Test Cases: {total_tcs}")
    total_p.runs[0].bold = True

    doc.add_page_break()

    # ---- Modules & test cases ----
    for module in MODULES:
        # Module heading
        mh = doc.add_heading(f"{module['id']} — {module['name']}", level=1)
        for run in mh.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

        doc.add_paragraph(
            f"This section covers {len(module['cases'])} test case(s) for the {module['name']} module."
        ).runs[0].font.italic = True

        doc.add_paragraph()

        for tc in module["cases"]:
            # Sub-heading for each test case
            sh = doc.add_heading(f"{tc['id']} — {tc['title']}", level=2)
            for run in sh.runs:
                run.font.size = Pt(11)

            add_tc_table(doc, tc)

        doc.add_page_break()

    doc.save(output_path)
    print(f"[OK] Document saved to: {output_path}")
    print(f"     Modules: {len(MODULES)}")
    print(f"     Total test cases: {total_tcs}")


if __name__ == "__main__":
    import os
    output = os.path.join(os.path.dirname(__file__), "AspireAI_Test_Cases.docx")
    build_document(output)
